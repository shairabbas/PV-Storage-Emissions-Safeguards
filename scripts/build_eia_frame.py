"""Construct an explicitly hybrid ERCOT stress test from public BA inputs.

2024 records are accessed only with --final, after the prospective freeze.
Participant load/PV are scaled BA profiles, not measured community data.
"""
from pathlib import Path
import argparse
import json
import numpy as np
import pandas as pd
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[1]
DATA=ROOT/'data'

def final_rows():
    core=['Balancing Authority','UTC Time at End of Hour','Demand (MW)',
          'Net Generation (MW) from Coal','Net Generation (MW) from Natural Gas']
    parts=[]
    for half in ('Jan_Jun','Jul_Dec'):
        path=DATA/f'EIA930_BALANCE_2024_{half}.csv'
        headers=pd.read_csv(path,nrows=0).columns.tolist()
        solar=[c for c in headers if c.startswith('Net Generation (MW) from Solar')
               and '(Imputed)' not in c and '(Adjusted)' not in c]
        wind=[c for c in headers if c.startswith('Net Generation (MW) from Wind')
              and '(Imputed)' not in c and '(Adjusted)' not in c]
        if not solar or not wind:raise ValueError('Unrecognized EIA renewable categories')
        for chunk in pd.read_csv(path,usecols=core+solar+wind,chunksize=50000,low_memory=False):
            part=chunk[chunk['Balancing Authority']=='ERCO']
            if len(part):
                part=part.copy()
                part['solar_mw']=part[solar].fillna(0).sum(axis=1)
                part['wind_mw']=part[wind].fillna(0).sum(axis=1)
                parts.append(part[core+['solar_mw','wind_mw']])
    x=pd.concat(parts,ignore_index=True)
    x['utc_end']=pd.to_datetime(x['UTC Time at End of Hour'],utc=True)
    wb=load_workbook(DATA/'EIA930_ERCO_hourly_2023_2025.xlsx',read_only=True,data_only=True)
    it=wb['Published Hourly Data'].iter_rows(values_only=True)
    head=next(it); ti=head.index('UTC time'); ei=head.index('CO2 Emissions Generated')
    emission=[]
    for row in it:
        stamp=row[ti]
        if stamp is None:continue
        if stamp.year<2024:continue
        if stamp.year>2024:break
        emission.append((pd.Timestamp(stamp,tz='UTC'),row[ei]))
    emissions=pd.DataFrame(emission,columns=['utc_end','eia_estimated_co2'])
    x=x.merge(emissions,on='utc_end',validate='one_to_one',how='left')
    x=x.rename(columns={'Demand (MW)':'demand_mw',
        'Net Generation (MW) from Coal':'coal_mw',
        'Net Generation (MW) from Natural Gas':'gas_mw',
        })
    return x[['utc_end','demand_mw','coal_mw','gas_mw','solar_mw','wind_mw','eia_estimated_co2']]

def main():
    ap=argparse.ArgumentParser();ap.add_argument('--final',action='store_true')
    args=ap.parse_args()
    if args.final and not (ROOT/'audit/V2_FREEZE.json').exists():
        raise SystemExit('Final data require the prospective freeze')
    dev=pd.read_csv(DATA/'eia_erco_2022_2023.csv.gz')
    if args.final:
        data=pd.concat([dev,final_rows()],ignore_index=True)
        output=DATA/'eia_erco_2022_2024_frame_corrected.csv.gz'
    else:
        data=dev
        output=DATA/'eia_erco_2022_2023_frame_corrected.csv.gz'
    data.utc_end=pd.to_datetime(data.utc_end,utc=True)
    data=data[(data.utc_end>=pd.Timestamp('2022-01-02',tz='UTC')) &
              (data.utc_end<pd.Timestamp('2025-01-01' if args.final else '2024-01-01',tz='UTC'))]
    data=data.sort_values('utc_end').drop_duplicates('utc_end')
    if args.final:
        # The prior source-year CSV ends with seven UTC hours of 2024-01-01
        # without the workbook's 2024 emissions estimates. This entire day
        # was prospectively excluded from final evaluation.
        data=data[~data.utc_end.between(pd.Timestamp('2024-01-01',tz='UTC'),
                  pd.Timestamp('2024-01-01 23:00',tz='UTC'))]
    # Keep only complete UTC days. Seven late-2023 CSV rows dated 2024 were
    # inadvertently visible in the development extract; they are excluded.
    counts=data.groupby(data.utc_end.dt.floor('D')).utc_end.transform('size')
    data=data[counts==24].copy().reset_index(drop=True)
    allowed={pd.Timedelta(hours=1)} if not args.final else {
        pd.Timedelta(hours=1),pd.Timedelta(hours=25)}
    if not set(data.utc_end.diff().dropna().unique()).issubset(allowed):
        raise ValueError('EIA hourly sequence has missing or duplicate timestamps')
    if data.eia_estimated_co2.isna().any():
        raise ValueError('Missing EIA-estimated CO2 on retained days')
    model=json.loads((ROOT/'audit/eia_response_model.json').read_text())
    coef=model['coefficients'];r=(data.demand_mw-data.solar_mw-data.wind_mw)/10000
    mef=np.clip((coef[1]+2*coef[2]*r)/10000,0,1.5)
    native=.2*data.demand_mw
    result=pd.DataFrame({'timestamp_utc':data.utc_end.astype(str),
        'native':native,'price_buy':np.where(data.utc_end.dt.hour.between(17,21),.24,
            np.where(data.utc_end.dt.hour.between(8,16),.14,.08)),
        'price_sell':.035,'eia_baseline_kg':.2*data.eia_estimated_co2,
        'mef_eia_erco':mef,
        'eia_curvature':coef[2]/(.2*10000**2),
        'aef_eia_erco':data.eia_estimated_co2/data.demand_mw})
    # EIA CO2 is metric tonnes per hour; tonnes/MWh = kg/kWh numerically.
    # The local baseline scales regional demand and emissions by 0.2 kW/MW.
    if ((data.demand_mw<=0).any() or (data.eia_estimated_co2<0).any()
            or not result.aef_eia_erco.between(0,1.5).all()
            or not np.allclose(result.eia_baseline_kg/result.native,
                               result.aef_eia_erco,rtol=1e-12,atol=1e-12)):
        raise ValueError('EIA source-unit or scaled-baseline dimensional check failed')
    # Calendar service is specified using only the 2022 mean hourly demand.
    train=data[data.utc_end.dt.year==2022]
    calendar=train.groupby(train.utc_end.dt.hour).demand_mw.mean()
    hour=data.utc_end.dt.hour.map(calendar).to_numpy()
    for i,(demand_share,solar_share) in enumerate(zip((.0024,.0019,.0015),(.025,.022,.018))):
        total=demand_share*data.demand_mw.to_numpy()
        flex0=.20*demand_share*hour
        result[f'critical_{i}']=total-flex0
        result[f'flex0_{i}']=flex0
        result[f'pv_{i}']=solar_share*data.solar_mw.to_numpy()
    q0=sum(result[f'critical_{i}']+result[f'flex0_{i}']-result[f'pv_{i}'] for i in range(3))
    result['background_eia_erco']=result.native-q0
    if (result[[f'critical_{i}' for i in range(3)]]<0).any().any():
        raise ValueError('Derived critical service is negative')
    result.to_csv(output,index=False,compression='gzip')
    print(json.dumps({'file':str(output.relative_to(ROOT)),'hours':len(result),
        'period':[result.timestamp_utc.iloc[0],result.timestamp_utc.iloc[-1]],
        'reported_measured_inputs':'BA demand, solar and wind generation',
        'estimated_input':'EIA production CO2',
        'model_generated':'scaled participant profiles and associational response'}))

if __name__=='__main__':main()
