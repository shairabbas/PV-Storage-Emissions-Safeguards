"""Build a separate editable Highlights file; enforce Elsevier's generic limit."""
from pathlib import Path
from docx import Document
ROOT=Path(__file__).resolve().parents[1]
lines=[x.strip() for x in (ROOT/'HIGHLIGHTS.txt').read_text(encoding='utf-8').splitlines() if x.strip()]
assert 3<=len(lines)<=5
assert all(len(x)<=85 for x in lines),[(len(x),x) for x in lines]
doc=Document();doc.add_heading('Highlights',level=1)
for line in lines:doc.add_paragraph(line,style='List Bullet')
doc.save(ROOT/'Highlights.docx')
print({'highlights':len(lines),'character_counts':[len(x) for x in lines]})
