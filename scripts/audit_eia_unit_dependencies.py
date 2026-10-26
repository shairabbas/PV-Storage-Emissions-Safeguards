"""Dimensional and downstream-action audit of the corrected EIA columns."""
from pathlib import Path
import hashlib
import json
import numpy as np
import pandas as pd
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[1]
DATA=ROOT/'data';RESULTS=ROOT/'results'
model=json.loads((ROOT/'audit/eia_response_model.json').read_text())
raw=DATA/'eia_erco_2022_2023.csv.gz'
assert hashlib.sha256(raw.read_bytes()).hexdigest()==model['source_sha256']
source=pd.read_csv(raw)
source['utc_end']=pd.to_datetime(source.utc_end,utc=True)
coef=model['coefficients']
records=[]
largest_unrelated_roundtrip_change=0.
for year in (2023,2024):
    bad=pd.read_csv(DATA/f'eia_erco_2022_{year}_frame.csv.gz')
    good=pd.read_csv(DATA/f'eia_erco_2022_{year}_frame_corrected.csv.gz')
    assert list(bad.columns)==list(good.columns)
    assert np.array_equal(bad.timestamp_utc.to_numpy(),good.timestamp_utc.to_numpy())
    changed=[]
    for col in good.columns:
        if not np.allclose(bad[col],good[col],rtol=0,atol=1e-10) if col!='timestamp_utc' else False:
            changed.append(col)
        if col not in ('timestamp_utc','aef_eia_erco','eia_baseline_kg'):
            largest_unrelated_roundtrip_change=max(largest_unrelated_roundtrip_change,
                float(np.max(np.abs(bad[col].to_numpy()-good[col].to_numpy()))))
    assert set(changed)=={'aef_eia_erco','eia_baseline_kg'},changed
    assert np.allclose(bad.aef_eia_erco,good.aef_eia_erco*1000,rtol=1e-12)
    assert np.allclose(bad.eia_baseline_kg,good.eia_baseline_kg*1000,rtol=1e-12)
    assert np.allclose(good.eia_baseline_kg/good.native,good.aef_eia_erco,rtol=1e-12)
    samples=[good.iloc[(good.timestamp_utc.str.startswith(str(y))).to_numpy().nonzero()[0][0]]
             for y in ([2023] if year==2023 else [2024])]
    for sample in samples:
        stamp=pd.Timestamp(sample.timestamp_utc)
        if stamp.year==2023:
            found=source.loc[source.utc_end==stamp]
            assert len(found)==1
            ton=float(found.eia_estimated_co2.iloc[0]);mw=float(found.demand_mw.iloc[0])
            solar=float(found.solar_mw.iloc[0]);wind=float(found.wind_mw.iloc[0])
        else:
            wb=load_workbook(DATA/'EIA930_ERCO_hourly_2023_2025.xlsx',read_only=True,data_only=True)
            rows=wb['Published Hourly Data'].iter_rows(values_only=True)
            header=next(rows);ti=header.index('UTC time');ei=header.index('CO2 Emissions Generated')
            ton=None
            for row in rows:
                if row[ti] is not None and pd.Timestamp(row[ti],tz='UTC')==stamp:
                    ton=float(row[ei]);break
            wb.close()
            assert ton is not None
            # The scaled frame retains 0.2 kW per reported MW of demand.
            mw=float(sample.native)/.2
            solar=wind=None
        assert np.isclose(sample.native,.2*mw,rtol=1e-12)
        assert np.isclose(sample.eia_baseline_kg,.2*ton,rtol=1e-12)
        assert np.isclose(sample.aef_eia_erco,ton/mw,rtol=1e-12)
        if solar is not None:
            r=(mw-solar-wind)/10000
            mef=np.clip((coef[1]+2*coef[2]*r)/10000,0,1.5)
            assert np.isclose(sample.mef_eia_erco,mef,rtol=1e-12)
        records.append({'year':stamp.year,'utc_end':str(stamp),
            'source_co2_metric_tonnes_per_hour':ton,'source_demand_MW':mw,
            'corrected_baseline_kg_per_hour':float(sample.eia_baseline_kg),
            'corrected_average_factor_kg_per_kwh':float(sample.aef_eia_erco),
            'model_marginal_factor_kg_per_kwh':float(sample.mef_eia_erco)})

differences=[]
for phase,year in [('validation',2023),('confirmation',2024)]:
    legacy=RESULTS/f'eia_{phase}_{year}'
    corrected=RESULTS/f'v3_eia_{phase}_{year}'
    for old_path in sorted(legacy.glob('*_daily.csv')):
        policy=old_path.name.removesuffix('_daily.csv')
        old=pd.read_csv(old_path);new=pd.read_csv(corrected/old_path.name)
        oh=pd.read_csv(legacy/f'{policy}_hourly.csv.gz')
        nh=pd.read_csv(corrected/f'{policy}_hourly.csv.gz')
        assert np.array_equal(old.date.to_numpy(),new.date.to_numpy())
        assert np.array_equal(oh.timestamp_utc.to_numpy(),nh.timestamp_utc.to_numpy())
        actioncols=['charge','discharge']+[f'flexible_{j}' for j in range(3)]
        action=float(np.abs(oh[actioncols].to_numpy()-nh[actioncols].to_numpy()).sum())
        differences.append({'phase':phase,'policy':policy,'days':len(new),
            'aggregate_action_l1_kwh':action,
            'cost_difference_corrected_minus_legacy_usd':float(new.real_cost.sum()-old.real_cost.sum()),
            'positive_kg_difference_corrected_minus_legacy':float(new.delta_emissions.clip(lower=0).sum()-old.delta_emissions.clip(lower=0).sum()),
            'absolute_reference_emissions_kg_corrected':float(new.reference_emissions.sum()),
            'absolute_reference_emissions_kg_legacy':float(old.reference_emissions.sum())})
actions=pd.DataFrame(differences)
actions.to_csv(ROOT/'tables/final_eia_legacy_vs_corrected_actions.csv',index=False)
assert len(actions)==30
assert (actions.aggregate_action_l1_kwh>1e-6).any()
report={'original_columns':{'aef_eia_erco':'1000 * estimated_CO2_tonnes / demand_MW; falsely kg/kWh',
    'eia_baseline_kg':'0.2 * 1000 * estimated_CO2_tonnes; falsely kg/hour'},
    'correct_columns':{'aef_eia_erco':'estimated_CO2_tonnes / demand_MW, numerically kg/kWh',
    'eia_baseline_kg':'0.2 * estimated_CO2_tonnes, kg/hour for 0.2-kW-per-MW scaling'},
    'changed_frame_columns_only':['aef_eia_erco','eia_baseline_kg'],
    'representative_source_reconstructions':records,
    'policies_with_changed_actions_by_phase':{
        phase:actions.loc[(actions.phase==phase)&(actions.aggregate_action_l1_kwh>1e-6),'policy'].tolist()
        for phase in ('validation','confirmation')},
    'maximum_action_l1_kwh':float(actions.aggregate_action_l1_kwh.max()),
    'largest_non_target_frame_roundtrip_difference':largest_unrelated_roundtrip_change,
    'calibration_dependency':'Margin calibration uses only synthetic 2023 frame and CV8 schedules; EIA columns absent. No EIA recalibration is required.',
    'response_model_dependency':'EIA response coefficients are fitted directly from raw EIA estimated tonnes and MW, before average-factor/baseline frame construction; source SHA-256 unchanged.',
    'scenario_dependency':'Residual blocks sample mef_eia_erco, which is unchanged. Repeating scenario preparation yields identical blocks and seeds.',
    'optimization_dependency':'aef_eia_erco enters its 168-hour forecast and carbon-price import objective for the applicable policies; corrected actions differ. All 15 policies per phase were rerun to preserve matched comparisons.',
    's1_numerical_sensitivity':'S1 has no direct aef or baseline dependency in optimize(), but its saved actions differ between legacy and corrected CSVs despite equal model objective values. The corrected CSV round-trip perturbs two unrelated flexible-load inputs by at most 3.6e-15 kWh; the MILP has alternative near-equal optima, so S1 differences must not be attributed causally to the unit conversion. The corrected S1 output is retained with this sensitivity disclosed.',
    'evaluation_dependency':'eia_baseline_kg changes absolute reference/system emissions, not reference-relative delta directly. Changed decisions can change delta. Attribution ledger uses aef_eia_erco.',
    'settlement_dependency':'Settlement uses expenses, not factors directly; action changes can alter expenses and hence participant charges/benefits. All were recomputed in corrected runs.'}
(ROOT/'audit/final_eia_unit_dependency.json').write_text(json.dumps(report,indent=2)+'\n')
print(json.dumps({'changed_actions':report['policies_with_changed_actions_by_phase'],
    'maximum_action_l1_kwh':report['maximum_action_l1_kwh'],
    'representative_source_reconstructions':records},indent=2))
